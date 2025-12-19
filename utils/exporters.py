"""
Export summaries to PDF and Word formats
"""
import io
from typing import List
from datetime import datetime


def export_to_pdf(
    original_text: str,
    extracted_sentences: List[str],
    final_summary: str,
    keywords: List[tuple] = None,
    method_used: str = "mmr"
) -> bytes:
    """
    Export summary to PDF with Arabic support.
    
    Returns:
        PDF file as bytes
    """
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    import arabic_reshaper
    from bidi.algorithm import get_display
    
    def reshape_arabic(text):
        """Reshape Arabic text for PDF rendering"""
        reshaped = arabic_reshaper.reshape(text)
        return get_display(reshaped)
    
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=50, leftMargin=50)
    
    styles = getSampleStyleSheet()
    
    # Arabic style (right-to-left)
    arabic_style = ParagraphStyle(
        'Arabic',
        parent=styles['Normal'],
        alignment=TA_RIGHT,
        fontSize=12,
        leading=18,
        wordWrap='RTL'
    )
    
    title_style = ParagraphStyle(
        'Title',
        parent=styles['Heading1'],
        alignment=TA_RIGHT,
        fontSize=16
    )
    
    story = []
    
    # Title
    story.append(Paragraph(reshape_arabic("تقرير التلخيص الآلي"), title_style))
    story.append(Spacer(1, 20))
    
    # Date
    date_str = datetime.now().strftime("%Y-%m-%d %H:%M")
    story.append(Paragraph(f"Date: {date_str} | Method: {method_used.upper()}", styles['Normal']))
    story.append(Spacer(1, 20))
    
    # Final Summary
    story.append(Paragraph(reshape_arabic("الملخص النهائي"), title_style))
    story.append(Spacer(1, 10))
    story.append(Paragraph(reshape_arabic(final_summary), arabic_style))
    story.append(Spacer(1, 20))
    
    # Extracted Sentences
    story.append(Paragraph(reshape_arabic("الجمل المستخرجة"), title_style))
    story.append(Spacer(1, 10))
    for i, sent in enumerate(extracted_sentences, 1):
        story.append(Paragraph(f"{i}. {reshape_arabic(sent)}", arabic_style))
        story.append(Spacer(1, 5))
    
    # Keywords
    if keywords:
        story.append(Spacer(1, 20))
        story.append(Paragraph(reshape_arabic("الكلمات المفتاحية"), title_style))
        story.append(Spacer(1, 10))
        kw_text = " | ".join([reshape_arabic(kw) for kw, _ in keywords])
        story.append(Paragraph(kw_text, arabic_style))
    
    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()


def export_to_docx(
    original_text: str,
    extracted_sentences: List[str],
    final_summary: str,
    keywords: List[tuple] = None,
    method_used: str = "mmr"
) -> bytes:
    """
    Export summary to Word document.
    
    Returns:
        DOCX file as bytes
    """
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    
    doc = Document()
    
    # Title
    title = doc.add_heading("تقرير التلخيص الآلي", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    
    # Date and method
    date_str = datetime.now().strftime("%Y-%m-%d %H:%M")
    p = doc.add_paragraph(f"Date: {date_str} | Method: {method_used.upper()}")
    
    # Final Summary
    doc.add_heading("الملخص النهائي", level=1).alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p = doc.add_paragraph(final_summary)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    
    # Extracted Sentences
    doc.add_heading("الجمل المستخرجة", level=1).alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for i, sent in enumerate(extracted_sentences, 1):
        p = doc.add_paragraph(f"{i}. {sent}")
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    
    # Keywords
    if keywords:
        doc.add_heading("الكلمات المفتاحية", level=1).alignment = WD_ALIGN_PARAGRAPH.RIGHT
        kw_text = " | ".join([kw for kw, _ in keywords])
        p = doc.add_paragraph(kw_text)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
